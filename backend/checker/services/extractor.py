"""Read plain text out of an uploaded resume (PDF, DOCX or TXT)."""
import io
import re

from docx import Document
from pypdf import PdfReader


class ResumeReadError(Exception):
    pass


ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt")


def _from_pdf(data):
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ResumeReadError("This PDF is password protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ResumeReadError:
        raise
    except Exception:
        raise ResumeReadError("We could not read this PDF. Please upload a valid PDF file.")
    return "\n".join(pages), len(reader.pages)


def _from_docx(data):
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise ResumeReadError("We could not read this DOCX file. Please upload a valid .docx file.")
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return "\n".join(lines), None


def extract_text(filename, data):
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text, pages = _from_pdf(data)
    elif name.endswith(".docx"):
        text, pages = _from_docx(data)
    elif name.endswith(".txt"):
        text, pages = data.decode("utf-8", errors="ignore"), None
    elif name.endswith(".doc"):
        raise ResumeReadError("Old .doc files are not supported. Please save your resume as PDF or DOCX and upload again.")
    else:
        raise ResumeReadError("Unsupported file type. Please upload a PDF, DOCX or TXT resume.")

    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise ResumeReadError(
            "We could not find enough text in this file. If it is a scanned image, "
            "please upload a text-based PDF or DOCX."
        )
    return text, pages
